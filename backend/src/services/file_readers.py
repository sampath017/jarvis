"""Streaming readers for Office, archives and opaque files; never execute sources."""
import codecs
import io
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import subprocess
import struct
import tempfile
import tarfile
import gzip
import bz2
import lzma
import zipfile
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

CATALOG_MIME = 'application/x-jarvis-catalog+json'
READER_VERSION = 2
TEXT_EXTENSIONS = {'.txt', '.csv', '.tsv', '.json', '.jsonl', '.ndjson', '.xml', '.yaml', '.yml', '.md', '.log', '.py', '.js', '.ts', '.java', '.kt', '.sql', '.sh', '.bat', '.ps1', '.ini', '.cfg', '.css', '.html', '.htm', '.svg', '.ipynb', '.eml', '.ics', '.vcf'}
OFFICE_EXTENSIONS = {'.doc', '.docx', '.docm', '.xls', '.xlsx', '.xlsm', '.xlsb', '.ppt', '.pptx', '.pptm', '.odt', '.ods', '.odp', '.rtf'}
ARCHIVE_EXTENSIONS = {'.zip', '.rar', '.7z', '.tar', '.gz', '.bz2', '.xz', '.tgz', '.epub', '.apk', '.jar', '.war'}


@dataclass
class Coverage:
    content: int = 0
    metadata: int = 0
    issues: Counter = field(default_factory=Counter)

    @property
    def level(self):
        return 'metadata' if not self.content else 'partial' if self.issues else 'content'

    def report(self):
        return {'coverage': self.level, 'issues': dict(self.issues), 'reader_version': READER_VERSION}


def text_units(text, locator):
    for offset in range(0, len(text), 3000):
        part = text[offset:offset + 3500]
        if part.strip():
            yield {'kind': 'text', 'text': part}, part, f'{locator}, text {offset}'


def catalog(state, title, reason, locator='file metadata', details='', attention=True):
    state.metadata += 1
    if attention:
        state.issues[reason] += 1
    for item, text, loc in text_units(f'File: {title}\nCoverage: metadata only. {reason}\n{details}', locator):
        item['_content_kind'] = 'metadata'
        yield item, text, loc


@contextmanager
def stream(raw):
    if isinstance(raw, (str, Path)):
        with open(raw, 'rb') as source:
            yield source
    else:
        yield io.BytesIO(raw)


@contextmanager
def materialize(source, suffix=''):
    # Large archive members use private storage instead of filling container RAM.
    if os.getenv('INDEX_STAGING_BUCKET') and Path('/sources').is_dir():
        from google.cloud import storage
        from uuid import uuid4
        name = 'reader-temp/' + uuid4().hex
        blob = storage.Client().bucket(os.environ['INDEX_STAGING_BUCKET']).blob(name)
        try:
            with blob.open('wb', chunk_size=8 * 1024 * 1024, ignore_flush=True) as target:
                shutil.copyfileobj(source, target, length=1024 * 1024)
            yield Path('/sources') / name
        finally:
            if blob.exists():
                blob.delete()
    else:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / ('source' + suffix)
            with target.open('wb') as output:
                shutil.copyfileobj(source, output, length=1024 * 1024)
            yield target


def xml_paragraphs(source):
    # Capture runs before clearing nodes, preserving words split by formatting.
    paragraph, stack = [], []
    for event, element in ET.iterparse(source, events=('start', 'end')):
        if event == 'start':
            stack.append(element)
            continue
        local = element.tag.rsplit('}', 1)[-1]
        if local == 't' and element.text:
            paragraph.append(element.text)
        elif local in ('tab', 'br'):
            paragraph.append(' ')
        elif local == 'p':
            if paragraph:
                yield ''.join(paragraph)
                paragraph = []
        stack.pop()
        if stack:
            stack[-1].remove(element)
        element.clear()
    if paragraph:
        yield ''.join(paragraph)


def office(raw, extension, title, state, basic):
    if extension in {'.doc', '.ppt', '.odt', '.ods', '.odp', '.rtf', '.xlsb'}:
        converted = '.xlsx' if extension in {'.ods', '.xlsb'} else '.pptx' if extension in {'.ppt', '.odp'} else '.docx'
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            profile = folder / 'profile'
            (profile / 'user').mkdir(parents=True)
            (profile / 'user/registrymodifications.xcu').write_text('''<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop><prop oor:name="SecureURL" oor:op="fuse"><value/></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>0</value></prop></item></oor:items>''')
            source = folder / ('input' + extension)
            if isinstance(raw, (str, Path)):
                source.symlink_to(Path(raw).resolve())
            else:
                source.write_bytes(raw)
            subprocess.run(['libreoffice', '-env:UserInstallation=' + profile.as_uri(), '--headless', '--norestore', '--nodefault',
                '--convert-to', converted[1:], '--outdir', str(folder), str(source)], check=True, capture_output=True,
                env={'PATH': os.getenv('PATH', ''), 'HOME': str(folder), 'LANG': 'C.UTF-8'})
            target = source.with_suffix(converted)
            if not target.exists():
                raise ValueError('Office conversion needs an unencrypted, readable document')
            yield from office(target, converted, title, state, basic)
        return
    if extension == '.xls':
        import xlrd
        from xlrd.formula import decompile_formula, FMLA_TYPE_CELL
        if isinstance(raw, (str, Path)):
            workbook = xlrd.open_workbook(filename=str(raw), on_demand=True)
        else:
            workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
        try:
            for name in workbook.sheet_names():
                formulas = {}
                original = workbook.get_record_parts
                def records():
                    record = original()
                    code, length, data = record
                    if code in {0x6, 0x206, 0x406} and len(data) >= 22:
                        row, column = struct.unpack('<HH', data[:4])
                        try:
                            expression = decompile_formula(workbook, data[22:], struct.unpack('<H', data[20:22])[0], FMLA_TYPE_CELL, browx=row, bcolx=column)
                        except Exception:
                            expression = None
                        formulas[row, column] = expression
                    return record
                workbook.get_record_parts = records
                try:
                    sheet = workbook.sheet_by_name(name)
                finally:
                    workbook.get_record_parts = original
                for row in range(sheet.nrows):
                    values = []
                    for col in range(sheet.ncols):
                        value = sheet.cell_value(row, col)
                        if sheet.cell_type(row, col) == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                        if value != '':
                            values.append(f'column {col + 1}: {value}')
                        if (row, col) in formulas:
                            expression = formulas[row, col]
                            if expression is None:
                                yield from catalog(state, title, 'A legacy/shared formula expression could not be decoded; cached cell values are included', f'sheet {name}, row {row + 1}, column {col + 1}')
                            else:
                                values.append(f'column {col + 1} formula: ={expression}')
                    yield from text_units(' | '.join(values), f'sheet {name}, row {row + 1}')
                workbook.unload_sheet(name)
        finally:
            workbook.release_resources()
        return
    if extension in {'.xlsx', '.xlsm'}:
        import openpyxl
        from itertools import zip_longest
        with stream(raw) as source, stream(raw) as cache_source:
            workbook = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
            cached = openpyxl.load_workbook(cache_source, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook:
                    yield from text_units(f'Sheet: {sheet.title}. Visibility: {sheet.sheet_state}', f'sheet {sheet.title}')
                    for row_number, (row, values) in enumerate(zip_longest(sheet.iter_rows(), cached[sheet.title].iter_rows(), fillvalue=()), 1):
                        cells = []
                        for cell, value in zip_longest(row, values):
                            if cell is not None and cell.value is not None:
                                content = f'{cell.coordinate}: {cell.value}'
                                if cell.number_format != 'General':
                                    content += f' (format: {cell.number_format})'
                                if cell.data_type == 'f' and value is not None and value.value is not None:
                                    content += f' (cached value: {value.value})'
                                cells.append(content)
                        if cells:
                            yield from text_units(' | '.join(cells), f'sheet {sheet.title}, row {row_number}')
            finally:
                workbook.close(); cached.close()
    elif extension in {'.pptx', '.pptm'}:
        from pptx import Presentation
        with stream(raw) as source:
            presentation = Presentation(source)
        def shapes(items):
            for shape in items:
                if hasattr(shape, 'shapes'):
                    yield from shapes(shape.shapes)
                if shape.has_text_frame:
                    yield shape.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield ' | '.join(cell.text for cell in row.cells)
                if shape.has_chart:
                    for series in shape.chart.series:
                        yield f'Chart series: {series.name}; values: {list(series.values)}'
        for n, slide in enumerate(presentation.slides, 1):
            for text in shapes(slide.shapes):
                yield from text_units(text, f'slide {n}')
            if slide.has_notes_slide:
                for text in shapes(slide.notes_slide.shapes):
                    yield from text_units(text, f'slide {n}, speaker notes')
    elif extension in {'.docx', '.docm'}:
        with stream(raw) as source, zipfile.ZipFile(source) as document:
            for name in document.namelist():
                if re.match(r'word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$', name):
                    with document.open(name) as body:
                        for n, paragraph in enumerate(xml_paragraphs(body), 1):
                            yield from text_units(paragraph, f'{name}, paragraph {n}')
    # Embedded images and chart labels remain searchable alongside document text.
    with stream(raw) as source, zipfile.ZipFile(source) as document:
        for name in document.namelist():
            if ('/media/' in name or '/embeddings/' in name or name.endswith('vbaProject.bin')) and not name.endswith('/'):
                with document.open(name) as member, materialize(member, Path(name).suffix) as image:
                    yield from read(image, mimetypes.guess_type(name)[0] or 'application/octet-stream', name, state, basic, prefix=title + ' / ')
            elif '/charts/' in name and name.endswith('.xml'):
                with document.open(name) as member:
                    values = []
                    for event, element in ET.iterparse(member, events=('end',)):
                        if element.tag.rsplit('}', 1)[-1] in {'v', 't'} and element.text:
                            values.append(element.text)
                        element.clear()
                    yield from text_units(' | '.join(values), f'chart {name}')
            elif (name.startswith('docProps/') or '/comments' in name or name.startswith('customXml/')) and name.endswith('.xml'):
                with document.open(name) as member:
                    for event, element in ET.iterparse(member, events=('end',)):
                        if element.text and element.text.strip():
                            yield from text_units(f"{element.tag.rsplit('}', 1)[-1]}: {element.text}", f'document metadata/comments {name}')
                        element.clear()


class Blocks(io.RawIOBase):
    def __init__(self, blocks):
        self.blocks, self.pending = iter(blocks), memoryview(b'')
    def readable(self): return True
    def readinto(self, target):
        if not self.pending:
            self.pending = memoryview(next(self.blocks, b''))
        count = min(len(target), len(self.pending))
        target[:count] = self.pending[:count]
        self.pending = self.pending[count:]
        return count


def archives(raw, title, state, basic, prefix):
    extension = Path(title).suffix.lower()
    if extension in {'.gz', '.bz2', '.xz'}:
        decoder = {'.gz': gzip.GzipFile, '.bz2': bz2.BZ2File, '.xz': lzma.LZMAFile}[extension]
        with stream(raw) as source:
            with (decoder(fileobj=source) if extension == '.gz' else decoder(source)) as body:
                with materialize(body) as path:
                    yield from read(path, mimetypes.guess_type(Path(title).stem)[0] or 'application/octet-stream', Path(title).stem, state, basic, prefix=prefix + title + ' / ')
        return
    if extension in {'.tar', '.tgz'}:
        with stream(raw) as source, tarfile.open(fileobj=source, mode='r|*') as archive:
            for member in archive:
                locator = prefix + title + ' / ' + member.name
                if member.isdir() or not member.isfile():
                    yield from catalog(state, member.name, 'Archive directory or link; no server filesystem target is followed', locator, str(member.linkname or ''), attention=False)
                    continue
                body = archive.extractfile(member)
                if Path(member.name).suffix.lower() in TEXT_EXTENSIONS:
                    yield from decoded_text(body, locator)
                else:
                    with materialize(body) as path:
                        yield from read(path, mimetypes.guess_type(member.name)[0] or 'application/octet-stream', member.name, state, basic, prefix=prefix + title + ' / ')
        return
    with stream(raw) as probe:
        zip_format = zipfile.is_zipfile(probe)
    if zip_format:
        with stream(raw) as source, zipfile.ZipFile(source) as archive:
            for member in archive.infolist():
                if member.is_dir():
                    yield from catalog(state, member.filename, 'Archive directory; child entries are indexed separately', prefix + title + ' / ' + member.filename, attention=False)
                    continue
                name = member.filename
                locator = prefix + title + ' / ' + name
                if member.flag_bits & 1:
                    yield from catalog(state, name, 'Encrypted archive member needs its password', locator)
                    continue
                try:
                    with archive.open(member) as body:
                        if Path(name).suffix.lower() in TEXT_EXTENSIONS and not body.peek(4).startswith(b'\x03\x00\x08\x00'):
                            yield from decoded_text(body, locator)
                        else:
                            with materialize(body, Path(name).suffix) as path:
                                yield from read(path, mimetypes.guess_type(name)[0] or 'application/octet-stream', name, state, basic, prefix=prefix + title + ' / ')
                except Exception:
                    yield from catalog(state, name, 'Archive member is encrypted, damaged, or could not be decoded', locator)
        return
    import libarchive
    with stream(raw) as source, libarchive.stream_reader(source) as archive:
        for member in archive:
            name = member.pathname
            locator = prefix + title + ' / ' + name
            if member.isdir:
                yield from catalog(state, name, 'Archive directory; child entries are indexed separately', locator, attention=False)
                continue
            if not member.isfile:
                yield from catalog(state, name, 'Archive link or special entry; target is not followed on the server', locator, str(member.linkpath or ''), attention=False)
                continue
            body = io.BufferedReader(Blocks(member.get_blocks()), buffer_size=64 * 1024)
            extension = Path(name).suffix.lower()
            try:
                if extension in TEXT_EXTENSIONS:
                    yield from decoded_text(body, locator)
                else:
                    with materialize(body, extension) as path:
                        yield from read(path, mimetypes.guess_type(name)[0] or 'application/octet-stream', name, state, basic, prefix=prefix + title + ' / ')
            except Exception:
                yield from catalog(state, name, 'Archive member is encrypted, damaged, or could not be decoded', locator)


def decoded_text(source, locator):
    probe = source.read(4)
    encoding = 'utf-16' if probe.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
    decoder = codecs.getincrementaldecoder(encoding)(errors='replace')
    buffer, offset = decoder.decode(probe), 0
    while True:
        part = source.read(64 * 1024)
        buffer += decoder.decode(part, final=not part)
        while len(buffer) >= 3500:
            text = buffer[:3500]
            if text.strip():
                yield {'kind': 'text', 'text': text}, text, f'{locator}, characters {offset + 1}-{offset + len(text)}'
            buffer, offset = buffer[3000:], offset + 3000
        if not part:
            if buffer.strip():
                yield from text_units(buffer, f'{locator}, characters {offset + 1}')
            break


def binary(raw, title, state, prefix):
    details = ''
    if Path(title).suffix.lower() in {'.exe', '.dll', '.sys'}:
        try:
            import pefile
            pe = pefile.PE(str(raw), fast_load=True) if isinstance(raw, (str, Path)) else pefile.PE(data=raw, fast_load=True)
            try:
                pe.parse_data_directories(directories=[pefile.DIRECTORY_ENTRY['IMAGE_DIRECTORY_ENTRY_RESOURCE']])
                info = {'architecture': pefile.MACHINE_TYPE.get(pe.FILE_HEADER.Machine), 'timestamp': pe.FILE_HEADER.TimeDateStamp}
                for groups in getattr(pe, 'FileInfo', []):
                    for group in groups:
                        for table in getattr(group, 'StringTable', []):
                            for key, value in table.entries.items():
                                info[key.decode(errors='replace')] = value.decode(errors='replace')
                details = json.dumps(info, ensure_ascii=False)
            finally:
                pe.close()
        except Exception:
            details = 'Executable header could not be interpreted; readable strings are still extracted.'
    yield from catalog(state, title, 'Binary format: extracted strings and metadata do not describe all program behavior', prefix + 'binary metadata', details)
    with stream(raw) as source:
        carry, offset, buffer, location = b'', 0, '', 0
        while part := source.read(64 * 1024):
            data = carry + part
            cutoff = max(0, len(data) - 1024)
            for match in re.finditer(rb'[\x20-\x7e]{6,}|(?:[\x20-\x7e]\x00){6,}', data):
                if match.start() >= cutoff:
                    continue
                value = data[match.start():min(match.end(), cutoff)]
                if b'\x00' in value:
                    value = value[:len(value) // 2 * 2]
                text = value.decode('utf-16-le' if b'\x00' in value else 'ascii')
                if not buffer:
                    location = max(0, offset - len(carry)) + match.start()
                buffer += text + '\n'
                while len(buffer) >= 3500:
                    yield from text_units(buffer[:3500], f'{prefix}binary strings near byte {location}')
                    buffer = buffer[3500:]
            carry = data[cutoff:]
            offset += len(part)
        for match in re.finditer(rb'[\x20-\x7e]{6,}|(?:[\x20-\x7e]\x00){6,}', carry):
            text = match.group().decode('utf-16-le' if b'\x00' in match.group() else 'ascii')
            buffer += text + '\n'
        yield from text_units(buffer, f'{prefix}binary strings near byte {location}')


def database(raw, title, state, basic, prefix):
    if not isinstance(raw, (str, Path)):
        with materialize(io.BytesIO(raw), '.sqlite') as path:
            yield from database(path, title, state, basic, prefix)
        return
    connection = sqlite3.connect(Path(raw).resolve().as_uri() + '?mode=ro&immutable=1', uri=True)
    try:
        connection.execute('PRAGMA query_only=ON')
        connection.execute('PRAGMA cache_size=-16384')
        for name, schema in connection.execute("SELECT name, sql FROM sqlite_master WHERE type='table'"):
            yield from text_units(f'Table {name}: {schema}', f'{prefix}database schema')
            identifier = '"' + name.replace('"', '""') + '"'
            cursor = connection.execute('SELECT * FROM ' + identifier)
            columns = [d[0] for d in cursor.description]
            for n, row in enumerate(cursor, 1):
                text = []
                for column, value in zip(columns, row):
                    if isinstance(value, bytes):
                        yield from read(value, 'application/octet-stream', column, state, basic, prefix=f'{prefix}table {name}, row {n}, blob ')
                    else:
                        text.append(f'{column}: {value}')
                yield from text_units(' | '.join(text), f'{prefix}table {name}, row {n}')
    finally:
        connection.close()


def read(raw, mime, title, state, basic, prefix=''):
    before = state.content + state.metadata
    extension = Path(title).suffix.lower()
    extension = {
        'application/msword': '.doc', 'application/vnd.ms-excel': '.xls', 'application/vnd.ms-powerpoint': '.ppt',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document': '.docx',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': '.xlsx',
        'application/vnd.openxmlformats-officedocument.presentationml.presentation': '.pptx',
    }.get(mime, extension)
    with stream(raw) as source:
        magic = source.read(4096)
    if magic.startswith(b'%PDF'):
        mime = 'application/pdf'
    elif mime == 'application/octet-stream':
        mime = mimetypes.guess_type(title)[0] or mime
    if not magic:
        yield from catalog(state, title, 'Empty file; filename and source metadata are available', prefix + 'file metadata')
        return
    try:
        if mime == CATALOG_MIME:
            with stream(raw) as source:
                data = json.load(source)
            yield from catalog(state, title, data.get('reason', 'Source content not available'), prefix + 'Drive metadata', json.dumps(data.get('memory', {}), ensure_ascii=False), attention=data.get('needs_attention', True))
            return
        if mime == 'application/x-jarvis-form+ndjson':
            with stream(raw) as source:
                for n, line in enumerate(source, 1):
                    data = json.loads(line)
                    if data.get('kind') == 'reader_issue':
                        yield from catalog(state, title, data['reason'], f'{prefix}Forms access report')
                    else:
                        for unit in text_units(json.dumps(data, ensure_ascii=False), f'{prefix}form record {n}'):
                            state.content += 1
                            yield unit
            return
        if magic.startswith(b'\x03\x00\x08\x00'):
            from pyaxmlparser.axmlprinter import AXMLPrinter
            with stream(raw) as source:
                printer = AXMLPrinter(source.read())
            if not printer.is_valid():
                raise ValueError('Invalid Android binary XML')
            for unit in text_units(printer.get_xml().decode('utf-8'), prefix + 'Android manifest/resource XML'):
                state.content += 1
                yield unit
            return
        if magic.startswith(b'SQLite format 3\x00'):
            iterator = database(raw, title, state, basic, prefix)
        elif extension in OFFICE_EXTENSIONS:
            iterator = office(raw, extension, title, state, basic)
        elif extension in ARCHIVE_EXTENSIONS or mime in {'application/zip', 'application/x-7z-compressed', 'application/vnd.rar', 'application/x-rar-compressed', 'application/x-tar'} or magic.startswith(b'PK\x03\x04'):
            iterator = archives(raw, title, state, basic, prefix)
        elif mime.startswith('text/') and not magic.startswith((b'\xff\xfe', b'\xfe\xff')):
            iterator = basic(raw, mime, title)
        elif extension in TEXT_EXTENSIONS or mime in {'application/json', 'application/xml', 'application/javascript', 'application/x-ndjson', 'application/x-jarvis-form+ndjson'} or magic.startswith((b'\xff\xfe', b'\xfe\xff')):
            with stream(raw) as source:
                for unit in decoded_text(source, prefix + title):
                    state.content += 1
                    yield unit
            return
        elif mime == 'application/pdf' or mime.startswith(('image/', 'audio/', 'video/')):
            iterator = basic(raw, mime, title)
        else:
            # Unknown uploads often have an octet-stream MIME type despite being text.
            if magic and b'\x00' not in magic:
                try:
                    magic.decode('utf-8')
                    with stream(raw) as source:
                        for unit in decoded_text(source, prefix + title):
                            state.content += 1
                            yield unit
                    return
                except UnicodeError:
                    pass
            iterator = binary(raw, title, state, prefix)
        for item, text, locator in iterator:
            if item.get('_content_kind') != 'metadata':
                state.content += 1
            yield item, text, prefix + locator
    except Exception as exc:
        # Keep the file in the index, with honest coverage and a visible reason.
        reason = 'Content needs a password, a working decoder, or an intact source'
        if isinstance(exc, FileNotFoundError):
            reason = 'Required document reader is unavailable'
        yield from catalog(state, title, reason, prefix + 'reader report')
    if state.content + state.metadata == before:
        yield from catalog(state, title, 'Empty file or no readable content', prefix + 'reader report')

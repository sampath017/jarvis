"""Container-only tests: legacy Office conversion, RAR/7z, and actual APK XML."""
import io
import subprocess
import tempfile
import zipfile
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from src.services.drive_index import chunks
from src.services.file_readers import Coverage

def read(path, mime):
    state = Coverage()
    text = '\n'.join(unit[1] for unit in chunks(path, mime, path.name, state))
    assert state.content, (path.name, state.report())
    return text

with tempfile.TemporaryDirectory() as folder:
    folder = Path(folder)
    doc = Document(); doc.add_paragraph('Legacy writer calibration ORBIT-417'); doc.save(folder/'writer.docx')
    book = Workbook(); book.active['A1']='Legacy workbook ORBIT-417'; book.active['B1']='=SUM(1,2)'; book.save(folder/'workbook.xlsx')
    deck = Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text='Legacy slides ORBIT-417'; deck.save(folder/'slides.pptx')
    for source, output, mime in [('writer.docx','doc','application/msword'),('workbook.xlsx','xls','application/vnd.ms-excel'),('slides.pptx','ppt','application/vnd.ms-powerpoint')]:
        subprocess.run(['libreoffice','--headless','--convert-to',output,'--outdir',str(folder),str(folder/source)],check=True,capture_output=True)
        target=(folder/source).with_suffix('.'+output)
        text = read(target,mime)
        assert 'ORBIT-417' in text
        if output == 'xls': assert 'formula: =SUM' in text, text
        print('Legacy '+output+' extraction passed',flush=True)
    (folder/'note.txt').write_text('Seven-zip calibration ORBIT-417')
    subprocess.run(['bsdtar','--format=7zip','-cf',str(folder/'archive.7z'),'-C',str(folder),'note.txt'],check=True)
    assert 'ORBIT-417' in read(folder/'archive.7z','application/x-7z-compressed')
    print('7z streaming extraction passed',flush=True)
    assert read(Path(__file__).with_name('libarchive.rar'),'application/vnd.rar')
    print('Official libarchive RAR fixture extraction passed',flush=True)
    manifest = Path(__file__).with_name('android-manifest.xml')
    if manifest.exists():
        text = read(manifest, 'application/xml')
        assert 'com.example.synthetic' in text and '42' in text and '1.0' in text
        print('Android manifest package/version extraction passed',flush=True)

import io
import json
import sqlite3
import zipfile
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from src.services.drive_index import chunks
from src.services.file_readers import Coverage


def passages(path, mime):
    coverage = Coverage()
    units = list(chunks(path, mime, path.name, coverage))
    return '\n'.join(u[1] for u in units), [u[2] for u in units], coverage


def test_word_preserves_tables_headers_and_formatted_words(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('Calibration ').bold = True
    p.add_run('ORBIT-417')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Inventory 42'
    doc.sections[0].header.paragraphs[0].text = 'Private header note'
    path = tmp_path / 'source.docx'; doc.save(path)
    text, locators, state = passages(path, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    assert all(value in text for value in ['Calibration ORBIT-417', 'Inventory 42', 'Private header note'])
    assert any('paragraph' in l for l in locators)
    assert state.level == 'content'


def test_spreadsheet_includes_hidden_sheet_and_formulas(tmp_path):
    book = Workbook(); book.active['A1'] = 'ORBIT-417'; book.active['B1'] = '=SUM(1,2)'
    hidden = book.create_sheet('Hidden inventory'); hidden.sheet_state = 'hidden'; hidden['B5'] = 'Secret stock 42'
    path = tmp_path / 'source.xlsx'; book.save(path)
    text, locators, state = passages(path, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert all(value in text for value in ['ORBIT-417', '=SUM(1,2)', 'Secret stock 42', 'Visibility: hidden'])
    assert any('Hidden inventory' in l for l in locators)
    assert state.level == 'content'


def test_presentation_reads_slide_and_speaker_notes(tmp_path):
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = 'ORBIT-417 launch'
    slide.notes_slide.notes_text_frame.text = 'Speaker note: delay until Friday'
    path = tmp_path / 'source.pptx'; presentation.save(path)
    text, locators, state = passages(path, 'application/vnd.openxmlformats-officedocument.presentationml.presentation')
    assert 'ORBIT-417 launch' in text and 'delay until Friday' in text
    assert any('speaker notes' in l for l in locators)
    assert state.level == 'content'


def test_nested_archive_reads_all_members_without_writing_member_paths(tmp_path):
    nested = io.BytesIO()
    with zipfile.ZipFile(nested, 'w') as archive: archive.writestr('inner.txt', 'Nested launch ORBIT-417')
    path = tmp_path / 'source.zip'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('../outside.txt', 'Source path is treated as a locator')
        archive.writestr('nested.zip', nested.getvalue())
        archive.writestr('data.json', '{"inventory": 42}')
        archive.writestr('empty.file', b'')
    text, locators, state = passages(path, 'application/zip')
    assert 'ORBIT-417' in text and 'inventory' in text and 'Source path' in text
    assert not (tmp_path.parent / 'outside.txt').exists()
    assert any('inner.txt' in l for l in locators)
    assert 'Empty file' in text and state.level == 'partial'


def test_binary_strings_are_searchable_without_executing_source(tmp_path):
    path = tmp_path / 'source.exe'; path.write_bytes(b'MZ\x00\xffProduct name ORBIT-417\x00' + 'Version 42'.encode('utf-16-le'))
    text, locators, state = passages(path, 'application/octet-stream')
    assert 'ORBIT-417' in text and 'Version 42' in text
    assert state.level == 'partial'


def test_unknown_opaque_file_remains_catalogued(tmp_path):
    path = tmp_path / 'opaque.blob'; path.write_bytes(bytes(range(32)))
    text, locators, state = passages(path, 'application/octet-stream')
    assert 'opaque.blob' in text and state.level == 'metadata'


def test_sqlite_reads_schema_rows_and_embedded_content_read_only(tmp_path):
    path = tmp_path / 'source.db'
    with sqlite3.connect(path) as db:
        db.execute('CREATE TABLE inventory(name TEXT, stock INTEGER)')
        db.execute('INSERT INTO inventory VALUES (?,?)', ('ORBIT-417', 42))
    text, locators, state = passages(path, 'application/octet-stream')
    assert 'inventory' in text and 'ORBIT-417' in text and 'stock: 42' in text


def test_forms_reports_missing_response_permission_as_partial(tmp_path):
    path = tmp_path / 'form.ndjson'
    path.write_text(json.dumps({'kind': 'form', 'body': {'title': 'ORBIT-417'}}) + '\n' + json.dumps({'kind': 'reader_issue', 'reason': 'Response access denied'}) + '\n')
    text, locators, state = passages(path, 'application/x-jarvis-form+ndjson')
    assert 'ORBIT-417' in text and 'Response access denied' in text
    assert state.level == 'partial'
